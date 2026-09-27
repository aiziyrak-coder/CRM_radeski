"""Operator daily report, KPIs and their time series (TZ 1.1, 4.11). All figures are computed
from source tables.

Filters (TZ 4.11 "davr, filial, operator, manba, xizmat yo'nalishi") apply where the data has
that dimension:
- branch, service category: appointment figures (bookings, visits, no-shows, confirmations,
  repeat visits); calls and inquiries have no branch or service;
- operator: calls, call results, QA scores and the appointments the operator created;
- source: inquiries and appointments (their "reklama manbasi").
"""

import re
import uuid
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from statistics import median
from typing import Any

from sqlalchemy import and_, distinct, exists, func, literal_column, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import clinic_time
from app.core.config import get_settings
from app.modules.ai.models import AnalysisStatus, CallAnalysis
from app.modules.catalog.models import Service
from app.modules.leads.models import Lead, LeadStage
from app.modules.patients.models import Source
from app.modules.scheduling.models import (
    Appointment,
    AppointmentService,
    AppointmentStatus,
    Recommendation,
    RecommendationStatus,
)
from app.modules.tasks.models import REACHED, Outcome, Task, TaskAttempt, TaskType
from app.modules.telephony.models import UNANSWERED_INBOUND, Call, CallDirection, CallStatus
from app.modules.users.models import Role, User

S = AppointmentStatus
# roles whose calls the reports break down per person (the daily report's operator filter)
CALLING_ROLES = (Role.OPERATOR, Role.SUPERVISOR)
# campaign-like calls whose patients count as "returned" once they visit (TZ 4.11)
RETURN_TYPES = (TaskType.CAMPAIGN, TaskType.REACTIVATION)
VISITED = (S.ARRIVED, S.COMPLETED)
# the KPI tiles compared with the previous period (dashboard deltas)
COMPARED = (
    "inbound_calls", "inbound_missed", "inbound_answer_rate", "outbound_calls", "leads_total",
    "leads_handled", "lead_to_booking", "first_response_median_min", "sla_breached", "attempts",
    "dial_rate", "confirmation_rate", "booking_to_visit", "no_show_rate", "repeat_rate",
    "returned_patients", "bookings", "visits", "qa_score", "missed_callback_avg_min",
    "missed_not_called_back",
)  # fmt: skip


@dataclass(frozen=True)
class Filters:
    branch_id: uuid.UUID | None = None
    user_id: uuid.UUID | None = None
    source: Source | None = None
    category_id: uuid.UUID | None = None

    def as_dict(self) -> dict[str, str | None]:
        return {k: str(v) if v is not None else None for k, v in self.__dict__.items()}


NO_FILTERS = Filters()


def _range(date_from: date, date_to: date) -> tuple[datetime, datetime]:
    return clinic_time.day_bounds(date_from)[0], clinic_time.day_bounds(date_to)[1]


def _pct(part: int, whole: int) -> float | None:
    return round(100 * part / whole, 1) if whole else None


def _minutes(seconds: float | None) -> float | None:
    return round(seconds / 60, 1) if seconds is not None else None


async def _count(session: AsyncSession, stmt) -> int:
    return await session.scalar(select(func.count()).select_from(stmt.subquery())) or 0


def _local_day(column):
    """The clinic-calendar date of a timestamptz column (for daily grouping). The zone is inlined
    (it's a setting, checked here) so SELECT and GROUP BY render the very same expression."""
    tz = get_settings().tz
    if not re.fullmatch(r"[A-Za-z_]+(/[A-Za-z_+-]+)*", tz):
        raise ValueError(f"bad TZ setting: {tz!r}")
    return func.date(func.timezone(literal_column(f"'{tz}'"), column))


def appointment_filters(f: Filters) -> list[Any]:
    conds: list[Any] = []
    if f.branch_id:
        conds.append(Appointment.branch_id == f.branch_id)
    if f.source:
        conds.append(Appointment.source == f.source)
    if f.user_id:
        conds.append(Appointment.created_by == f.user_id)
    if f.category_id:
        conds.append(
            exists().where(
                AppointmentService.appointment_id == Appointment.id,
                AppointmentService.service_id == Service.id,
                Service.category_id == f.category_id,
            )
        )
    return conds


def lead_filters(f: Filters) -> list[Any]:
    return [Lead.source == f.source] if f.source else []


def call_filters(f: Filters) -> list[Any]:
    return [Call.user_id == f.user_id] if f.user_id else []


async def call_stats(
    session: AsyncSession, start: datetime, end: datetime, user_id: uuid.UUID | None = None
) -> dict[str, Any]:
    """Telephony figures; all None while the PBX isn't connected (no call was ever logged)."""
    keys = (
        "inbound_calls", "inbound_answered", "inbound_missed", "inbound_answer_rate",
        "callbacks_requested", "avg_wait_sec", "outbound_calls", "outbound_answered",
        "talk_minutes",
    )  # fmt: skip
    if not await session.scalar(select(func.count()).select_from(Call)):
        return dict.fromkeys(keys)
    period = [Call.started_at >= start, Call.started_at < end]
    if user_id:  # an operator's own view: calls they handled
        period.append(Call.user_id == user_id)
    answered = Call.status == CallStatus.ANSWERED
    inbound = Call.direction == CallDirection.IN
    row = (
        await session.execute(
            select(
                func.count().filter(inbound),
                func.count().filter(inbound, answered),
                func.count().filter(inbound, Call.status.in_(UNANSWERED_INBOUND)),
                func.count().filter(inbound, Call.callback_requested),
                func.avg(Call.wait_seconds).filter(inbound, Call.wait_seconds.is_not(None)),
                func.count().filter(~inbound),
                func.count().filter(~inbound, answered),
                func.coalesce(func.sum(Call.talk_seconds), 0),
            ).where(*period)
        )
    ).one()
    total_in, in_answered, missed, callbacks, avg_wait, out, out_answered, talk = row
    return {
        "inbound_calls": total_in,
        "inbound_answered": in_answered,
        "inbound_missed": missed,
        "inbound_answer_rate": _pct(in_answered, total_in),
        "callbacks_requested": callbacks,
        "avg_wait_sec": round(float(avg_wait)) if avg_wait is not None else None,
        "outbound_calls": out,
        "outbound_answered": out_answered,
        "talk_minutes": round(talk / 60),
    }


def first_callback_at():
    """When a missed inbound call was called back: the first outbound call to the same number (a
    correlated subquery on `Call`; None = not called back within the configured window)."""
    window = timedelta(days=get_settings().missed_callback_window_days)
    back = Call.__table__.alias("callback")
    first_back = (
        select(func.min(back.c.started_at))
        .where(
            back.c.direction == CallDirection.OUT,
            back.c.phone == Call.phone,
            back.c.started_at > Call.started_at,
            back.c.started_at < Call.started_at + window,
        )
        .correlate(Call)
        .scalar_subquery()
    )
    return first_back


async def missed_callbacks(session: AsyncSession, start: datetime, end: datetime) -> dict[str, Any]:
    """TZ 4.11 "Javobsiz qo'ng'iroqlar: soni va ularga qayta qo'ng'iroq qilish vaqti"."""
    first_back = first_callback_at()
    rows = (
        await session.execute(
            select(Call.phone, func.extract("epoch", first_back - Call.started_at)).where(
                Call.direction == CallDirection.IN,
                Call.status.in_(UNANSWERED_INBOUND),
                Call.started_at >= start,
                Call.started_at < end,
            )
        )
    ).all()
    # anonymous callers (no number) can't be called back: they are not held against anyone
    delays = [float(d) for phone, d in rows if phone and d is not None]
    return {
        "missed_total": len(rows),
        "missed_called_back": len(delays),
        "missed_not_called_back": sum(1 for phone, d in rows if phone and d is None),
        "missed_callback_avg_min": _minutes(sum(delays) / len(delays)) if delays else None,
        "missed_callback_median_min": _minutes(median(delays)) if delays else None,
    }


async def qa_scores(
    session: AsyncSession, start: datetime, end: datetime, user_id: uuid.UUID | None = None
) -> dict[uuid.UUID | None, tuple[int, float | None]]:
    """(analysed calls, average AI score) per operator; the None key is everyone together."""
    stmt = (
        select(Call.user_id, func.count(), func.avg(CallAnalysis.score))
        .join(Call, Call.id == CallAnalysis.call_id)
        .where(
            CallAnalysis.status == AnalysisStatus.READY,
            Call.started_at >= start,
            Call.started_at < end,
        )
        .group_by(Call.user_id)
    )
    if user_id:
        stmt = stmt.where(Call.user_id == user_id)
    out: dict[uuid.UUID | None, tuple[int, float | None]] = {}
    total, weighted = 0, 0.0
    scored = 0
    for uid, n, avg in await session.execute(stmt):
        out[uid] = (n, round(float(avg)) if avg is not None else None)
        total += n
        if avg is not None:
            scored += n
            weighted += float(avg) * n
    out[None] = (total, round(weighted / scored) if scored else None)
    return out


async def daily(
    session: AsyncSession, day: date, user_id: uuid.UUID | None = None, f: Filters = NO_FILTERS
) -> dict[str, Any]:
    start, end = clinic_time.day_bounds(day)
    f = Filters(f.branch_id, user_id, f.source, f.category_id)
    attempts = select(TaskAttempt).where(
        TaskAttempt.created_at >= start, TaskAttempt.created_at < end
    )
    if user_id:
        attempts = attempts.where(TaskAttempt.user_id == user_id)
    # every result incl. the ones a booking wrote (booking figures) ...
    sub = attempts.subquery()
    # ... and only the operators' own calls (attempts, dial rate, outcomes)
    calls = attempts.where(TaskAttempt.automatic.is_(False)).subquery()

    by_outcome = dict(
        (
            await session.execute(select(calls.c.outcome, func.count()).group_by(calls.c.outcome))
        ).all()
    )
    reasons = dict(
        (
            await session.execute(
                select(calls.c.reason, func.count())
                .where(calls.c.reason.is_not(None))
                .group_by(calls.c.reason)
            )
        ).all()
    )
    campaign = dict(
        (
            await session.execute(
                select(sub.c.outcome, func.count())
                .where(sub.c.task_type.in_((TaskType.CAMPAIGN, TaskType.REACTIVATION)))
                .group_by(sub.c.outcome)
            )
        ).all()
    )
    repeat_bookings = await session.scalar(
        select(func.count())
        .select_from(sub)
        .where(
            sub.c.outcome == Outcome.BOOKED,
            sub.c.task_type.in_((TaskType.REPEAT_VISIT, TaskType.COURSE_CONTINUE)),
        )
    )
    appt = appointment_filters(f)
    booked = select(Appointment.id).where(
        Appointment.created_at >= start, Appointment.created_at < end, *appt
    )
    changed = and_(Appointment.status_changed_at >= start, Appointment.status_changed_at < end)
    no_shows = select(Appointment.id).where(
        Appointment.starts_at >= start,
        Appointment.starts_at < end,
        Appointment.status == S.NO_SHOW,
        *appt,
    )
    total_attempts = sum(by_outcome.values())
    reached = sum(n for o, n in by_outcome.items() if o in REACHED)
    return {
        "date": day.isoformat(),
        "filters": f.as_dict(),
        **await call_stats(session, start, end, user_id),
        "outbound_attempts": total_attempts,
        "reached": reached,
        "dial_rate": _pct(reached, total_attempts),
        "new_leads": await _count(
            session,
            select(Lead.id).where(
                Lead.created_at >= start, Lead.created_at < end, *lead_filters(f)
            ),
        ),
        "booked": await _count(session, booked),
        "repeat_bookings": repeat_bookings or 0,
        "not_booked": by_outcome.get(Outcome.REFUSED, 0) + by_outcome.get(Outcome.THINKING, 0),
        "reasons": {str(k): v for k, v in reasons.items()},
        "cancellations": await _count(
            session,
            select(Appointment.id).where(changed, Appointment.status == S.CANCELLED, *appt),
        ),
        "reschedules": await _count(
            session,
            select(Appointment.id).where(changed, Appointment.status == S.RESCHEDULED, *appt),
        ),
        "no_shows": await _count(session, no_shows),
        "outcomes": {str(k): v for k, v in by_outcome.items()},
        "campaign_outcomes": {str(k): v for k, v in campaign.items()},
    }


async def kpi(
    session: AsyncSession,
    date_from: date,
    date_to: date,
    f: Filters = NO_FILTERS,
    *,
    with_operators: bool = True,
) -> dict[str, Any]:
    start, end = _range(date_from, date_to)
    now = clinic_time.now()
    appt = appointment_filters(f)

    leads_in = (
        select(Lead)
        .where(Lead.created_at >= start, Lead.created_at < end, *lead_filters(f))
        .subquery()
    )
    leads_total = await session.scalar(select(func.count()).select_from(leads_in)) or 0
    leads_booked = (
        await session.scalar(
            select(func.count()).select_from(leads_in).where(leads_in.c.appointment_id.is_not(None))
        )
        or 0
    )
    median_seconds = await session.scalar(
        select(
            func.percentile_cont(0.5).within_group(
                func.extract("epoch", leads_in.c.first_response_at - leads_in.c.created_at)
            )
        ).where(leads_in.c.first_response_at.is_not(None))
    )
    breached = (
        await session.scalar(
            select(func.count())
            .select_from(leads_in)
            .where(func.coalesce(leads_in.c.first_response_at, now) > leads_in.c.sla_due_at)
        )
        or 0
    )
    # TZ 4.11 "Ishlov berilgan murojaatlar": inquiries somebody got in touch with in the period
    handled = await _count(
        session,
        select(Lead.id).where(
            Lead.first_response_at >= start, Lead.first_response_at < end, *lead_filters(f)
        ),
    )

    # operators' calls only: results written by a booking (automatic) are not dial attempts
    attempt_where = [
        TaskAttempt.created_at >= start,
        TaskAttempt.created_at < end,
        TaskAttempt.automatic.is_(False),
    ]
    if f.user_id:
        attempt_where.append(TaskAttempt.user_id == f.user_id)
    attempts = select(TaskAttempt).where(*attempt_where).subquery()
    total_attempts = await session.scalar(select(func.count()).select_from(attempts)) or 0
    reached = (
        await session.scalar(
            select(func.count()).select_from(attempts).where(attempts.c.outcome.in_(REACHED))
        )
        or 0
    )

    # TZ 4.11 "Tasdiqlash %": confirmed / the day's appointments. Days up to today only (future
    # visits aren't due for confirmation yet); a rescheduled visit lives on as its new booking.
    _, end_of_today = clinic_time.day_bounds(clinic_time.today())
    # a real confirmation: an operator's "confirmed" call result (a patient who simply walked in
    # closes the confirmation task too, but nobody confirmed anything)
    confirm_task = exists().where(
        Task.appointment_id == Appointment.id,
        Task.type == TaskType.CONFIRM_VISIT,
        TaskAttempt.task_id == Task.id,
        TaskAttempt.outcome == Outcome.CONFIRMED,
        TaskAttempt.automatic.is_(False),
    )
    day_appts = (
        select(
            Appointment.id,
            ((Appointment.status == S.CONFIRMED) | confirm_task).label("confirmed"),
        )
        .where(
            Appointment.starts_at >= start,
            Appointment.starts_at < min(end, end_of_today),
            Appointment.status != S.RESCHEDULED,
            *appt,
        )
        .subquery()
    )
    appts_of_day = await session.scalar(select(func.count()).select_from(day_appts)) or 0
    confirmed = (
        await session.scalar(
            select(func.count()).select_from(day_appts).where(day_appts.c.confirmed)
        )
        or 0
    )

    # TZ 4.11 "Qaytarilgan bemorlar": patients a campaign / reactivation call brought back who
    # actually came (a visit in the period, after the call task was created)
    returned = (
        await session.scalar(
            select(func.count(distinct(Appointment.patient_id))).where(
                Appointment.starts_at >= start,
                Appointment.starts_at < end,
                Appointment.status.in_(VISITED),
                exists().where(
                    Task.patient_id == Appointment.patient_id,
                    Task.type.in_(RETURN_TYPES),
                    Task.outcome == Outcome.BOOKED,
                    Task.created_at <= Appointment.starts_at,
                ),
                *appt,
            )
        )
        or 0
    )

    past = (
        select(Appointment.status)
        .where(
            Appointment.starts_at >= start,
            Appointment.starts_at < min(end, now),
            Appointment.status.in_((S.SCHEDULED, S.CONFIRMED, S.ARRIVED, S.COMPLETED, S.NO_SHOW)),
            *appt,
        )
        .subquery()
    )
    past_total = await session.scalar(select(func.count()).select_from(past)) or 0
    visited = (
        await session.scalar(
            select(func.count()).select_from(past).where(past.c.status.in_(VISITED))
        )
        or 0
    )
    no_show = (
        await session.scalar(
            select(func.count()).select_from(past).where(past.c.status == S.NO_SHOW)
        )
        or 0
    )
    bookings = await _count(
        session,
        select(Appointment.id).where(
            Appointment.created_at >= start, Appointment.created_at < end, *appt
        ),
    )

    rec_where = [Recommendation.due_date >= date_from, Recommendation.due_date <= date_to]
    if f.category_id:
        rec_where.append(
            exists().where(
                Service.id == Recommendation.service_id, Service.category_id == f.category_id
            )
        )
    if f.branch_id:
        rec_where.append(
            exists().where(
                Appointment.id == Recommendation.appointment_id,
                Appointment.branch_id == f.branch_id,
            )
        )
    recs = select(Recommendation.status).where(*rec_where).subquery()
    recs_total = await session.scalar(select(func.count()).select_from(recs)) or 0
    recs_booked = (
        await session.scalar(
            select(func.count())
            .select_from(recs)
            .where(recs.c.status == RecommendationStatus.BOOKED)
        )
        or 0
    )
    qa = await qa_scores(session, start, end, f.user_id)
    out: dict[str, Any] = {
        "from": date_from.isoformat(),
        "to": date_to.isoformat(),
        "filters": f.as_dict(),
        **await call_stats(session, start, end, f.user_id),
        **await missed_callbacks(session, start, end),
        "leads_total": leads_total,
        "leads_handled": handled,
        "lead_to_booking": _pct(leads_booked, leads_total),
        "first_response_median_min": _minutes(median_seconds),
        "sla_breached": breached,
        "attempts": total_attempts,
        "dial_rate": _pct(reached, total_attempts),
        "confirmation_rate": _pct(confirmed, appts_of_day),
        "bookings": bookings,
        "visits": visited,
        "booking_to_visit": _pct(visited, past_total),
        "no_show_rate": _pct(no_show, past_total),
        "repeat_rate": _pct(recs_booked, recs_total),
        "returned_patients": returned,
        "qa_analysed": qa[None][0],
        "qa_score": qa[None][1],
    }
    if with_operators:
        out["operators"] = await operator_board(session, start, end, attempts, f, qa)
    return out


async def operator_board(
    session: AsyncSession,
    start: datetime,
    end: datetime,
    attempts,
    f: Filters,
    qa: dict[uuid.UUID | None, tuple[int, float | None]],
) -> list[dict[str, Any]]:
    """Per operator: call results, bookings, talk time and QA score (leaderboard)."""
    rows = await session.execute(
        select(
            attempts.c.user_id,
            func.count(attempts.c.id),
            func.count().filter(attempts.c.outcome.in_(REACHED)),
            func.count().filter(attempts.c.outcome == Outcome.BOOKED),
        )
        .where(attempts.c.user_id.is_not(None))
        .group_by(attempts.c.user_id)
    )
    by_user: dict[uuid.UUID, dict[str, int]] = defaultdict(
        lambda: {"attempts": 0, "reached": 0, "booked_by_phone": 0}
    )
    for uid, n, r, b in rows:
        by_user[uid].update(attempts=n, reached=r, booked_by_phone=b)
    appt = appointment_filters(Filters(f.branch_id, None, f.source, f.category_id))
    bookings_by_user = dict(
        (
            await session.execute(
                select(Appointment.created_by, func.count())
                .where(
                    Appointment.created_at >= start,
                    Appointment.created_at < end,
                    Appointment.created_by.is_not(None),
                    *appt,
                )
                .group_by(Appointment.created_by)
            )
        ).all()
    )
    talk_by_user = dict(
        (
            await session.execute(
                select(Call.user_id, func.coalesce(func.sum(Call.talk_seconds), 0))
                .where(Call.started_at >= start, Call.started_at < end, Call.user_id.is_not(None))
                .group_by(Call.user_id)
            )
        ).all()
    )
    # whoever worked the phones: recorded call results, talked on the PBX or was scored by QA
    # (bookings alone don't put someone on the board: registrars book too)
    ids = set(by_user) | {u for u in talk_by_user if u} | {u for u in qa if u}
    names = dict(
        (await session.execute(select(User.id, User.full_name).where(User.id.in_(ids)))).all()
    )
    if f.user_id:
        names = {k: v for k, v in names.items() if k == f.user_id}
    board = []
    for uid, name in names.items():
        a = by_user[uid]
        qa_calls, qa_avg = qa.get(uid, (0, None))
        board.append(
            {
                "user_id": str(uid),
                "name": name,
                **a,
                "dial_rate": _pct(a["reached"], a["attempts"]),
                "appointments_created": bookings_by_user.get(uid, 0),
                "talk_minutes": round(talk_by_user.get(uid, 0) / 60),
                "qa_calls": qa_calls,
                "qa_score": qa_avg,
            }
        )
    board.sort(key=lambda o: (-o["appointments_created"], -o["attempts"], o["name"]))
    return board


async def kpi_compared(
    session: AsyncSession, date_from: date, date_to: date, f: Filters = NO_FILTERS
) -> dict[str, Any]:
    """KPIs of the period plus the same figures for the equally long period right before it."""
    current = await kpi(session, date_from, date_to, f)
    length = (date_to - date_from).days + 1
    prev_to = date_from - timedelta(days=1)
    prev_from = prev_to - timedelta(days=length - 1)
    previous = await kpi(session, prev_from, prev_to, f, with_operators=False)
    current["previous"] = {
        "from": prev_from.isoformat(),
        "to": prev_to.isoformat(),
        **{k: previous.get(k) for k in COMPARED},
    }
    return current


# --- time series ------------------------------------------------------------------------------


def _days(date_from: date, date_to: date) -> list[date]:
    return [date_from + timedelta(days=i) for i in range((date_to - date_from).days + 1)]


async def _per_day(session: AsyncSession, column, *where) -> dict[date, int]:
    day = _local_day(column)
    rows = await session.execute(select(day, func.count()).where(*where).group_by(day))
    return {d: n for d, n in rows}


async def series(
    session: AsyncSession, date_from: date, date_to: date, f: Filters = NO_FILTERS
) -> dict[str, Any]:
    """Daily figures for the trend charts, leads by source and the inquiry funnel."""
    start, end = _range(date_from, date_to)
    appt = appointment_filters(f)
    leads = await _per_day(
        session, Lead.created_at, Lead.created_at >= start, Lead.created_at < end, *lead_filters(f)
    )
    bookings = await _per_day(
        session,
        Appointment.created_at,
        Appointment.created_at >= start,
        Appointment.created_at < end,
        *appt,
    )
    visits = await _per_day(
        session,
        Appointment.starts_at,
        Appointment.starts_at >= start,
        Appointment.starts_at < end,
        Appointment.status.in_(VISITED),
        *appt,
    )
    no_shows = await _per_day(
        session,
        Appointment.starts_at,
        Appointment.starts_at >= start,
        Appointment.starts_at < end,
        Appointment.status == S.NO_SHOW,
        *appt,
    )
    call_day = _local_day(Call.started_at)
    inbound = Call.direction == CallDirection.IN
    calls: dict[date, tuple[int, int, int]] = {
        d: (i, o, m)
        for d, i, o, m in await session.execute(
            select(
                call_day,
                func.count().filter(inbound),
                func.count().filter(~inbound),
                func.count().filter(inbound, Call.status.in_(UNANSWERED_INBOUND)),
            )
            .where(Call.started_at >= start, Call.started_at < end, *call_filters(f))
            .group_by(call_day)
        )
    }
    qa: dict[date, tuple[int, float | None]] = {
        d: (n, round(float(avg)) if avg is not None else None)
        for d, n, avg in await session.execute(
            select(call_day, func.count(), func.avg(CallAnalysis.score))
            .join(Call, Call.id == CallAnalysis.call_id)
            .where(
                CallAnalysis.status == AnalysisStatus.READY,
                Call.started_at >= start,
                Call.started_at < end,
                *call_filters(f),
            )
            .group_by(call_day)
        )
    }
    days = []
    for d in _days(date_from, date_to):
        i, o, m = calls.get(d, (0, 0, 0))
        qa_n, qa_avg = qa.get(d, (0, None))
        days.append(
            {
                "date": d.isoformat(),
                "leads": leads.get(d, 0),
                "bookings": bookings.get(d, 0),
                "visits": visits.get(d, 0),
                "no_shows": no_shows.get(d, 0),
                "calls_in": i,
                "calls_out": o,
                "calls_missed": m,
                "qa_calls": qa_n,
                "qa_score": qa_avg,
            }
        )
    return {
        "from": date_from.isoformat(),
        "to": date_to.isoformat(),
        "filters": f.as_dict(),
        "days": days,
        "sources": await lead_sources(session, start, end, f),
        "funnel": await funnel(session, start, end, f),
    }


async def lead_sources(
    session: AsyncSession, start: datetime, end: datetime, f: Filters
) -> list[dict[str, Any]]:
    rows = await session.execute(
        select(Lead.source, func.count(), func.count().filter(Lead.appointment_id.is_not(None)))
        .where(Lead.created_at >= start, Lead.created_at < end, *lead_filters(f))
        .group_by(Lead.source)
    )
    out = [
        {"source": s.value if s else "unknown", "leads": n, "booked": b, "conversion": _pct(b, n)}
        for s, n, b in rows
    ]
    return sorted(out, key=lambda r: (-r["leads"], r["source"]))


async def funnel(
    session: AsyncSession, start: datetime, end: datetime, f: Filters
) -> list[dict[str, Any]]:
    """TZ 4.4 funnel of the period's inquiries: yangi -> aloqa -> yozildi -> tasdiqlandi ->
    keldi. Each stage includes the later ones, so the funnel never widens."""
    status = (
        select(Appointment.status)
        .where(Appointment.id == Lead.appointment_id)
        .correlate(Lead)
        .scalar_subquery()
    )
    visited = or_(Lead.stage == LeadStage.VISITED, status.in_(VISITED))
    confirmed = or_(visited, status == S.CONFIRMED)
    booked = or_(confirmed, Lead.appointment_id.is_not(None), Lead.stage == LeadStage.BOOKED)
    contacted = or_(booked, Lead.first_response_at.is_not(None), Lead.stage == LeadStage.CONTACTED)
    row = (
        await session.execute(
            select(
                func.count(),
                func.count().filter(contacted),
                func.count().filter(booked),
                func.count().filter(confirmed),
                func.count().filter(visited),
            ).where(Lead.created_at >= start, Lead.created_at < end, *lead_filters(f))
        )
    ).one()
    stages = ("new", "contacted", "booked", "confirmed", "visited")
    total = row[0]
    return [
        {"stage": s, "count": n, "share": _pct(n, total)} for s, n in zip(stages, row, strict=True)
    ]


# --- Excel ------------------------------------------------------------------------------------


def _fit(*sheets) -> None:
    for sheet in sheets:
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = max(
                12, *(len(str(c.value or "")) + 2 for c in column)
            )


def _save(wb) -> bytes:
    from io import BytesIO

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def kpi_workbook(data: dict[str, Any], labels: dict[str, str]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "KPI"
    ws.append(["Davr / Период", f"{data['from']} — {data['to']}"])
    for key, label in labels.items():
        ws.append([label, data.get(key)])
    ops = wb.create_sheet("Operatorlar")
    ops.append(
        [
            "Operator",
            "Urinishlar",
            "Gaplashildi",
            "Dozvon %",
            "Tel. orqali yozildi",
            "Yaratilgan qabullar",
            "Suhbat, daq",
            "QA ball",
        ]
    )
    for o in data["operators"]:
        ops.append(
            [
                o["name"],
                o["attempts"],
                o["reached"],
                o["dial_rate"],
                o["booked_by_phone"],
                o["appointments_created"],
                o["talk_minutes"],
                o["qa_score"],
            ]
        )
    _fit(ws, ops)
    return _save(wb)


def daily_workbook(data: dict[str, Any], labels: dict[str, str], who: str) -> bytes:
    """The operator's daily report (TZ 4.11) as a sheet the supervisor can file or send."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Kunlik hisobot"
    ws.append(["Sana / Дата", data["date"]])
    ws.append(["Operator", who])
    for key, label in labels.items():
        ws.append([label, data.get(key)])
    details = wb.create_sheet("Natijalar")
    details.append(["Bo'lim / Раздел", "Kod / Код", "Soni / Кол-во"])
    for section, key in (
        ("Natijalar / Результаты", "outcomes"),
        ("Sabablar / Причины", "reasons"),
        ("Kampaniya / Кампания", "campaign_outcomes"),
    ):
        for code, n in sorted(data[key].items()):
            details.append([section, code, n])
    _fit(ws, details)
    return _save(wb)


async def operators(session: AsyncSession) -> list[dict[str, Any]]:
    """Active call-center staff, for the report's operator filter."""
    rows = await session.execute(
        select(User.id, User.full_name)
        .where(User.is_active.is_(True), User.role.in_(CALLING_ROLES))
        .order_by(User.full_name)
    )
    return [{"id": str(uid), "full_name": name} for uid, name in rows]
