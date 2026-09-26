"""Pure parsing of the clinic's legacy Excel exports (no database access).

Sources (see docs/01_TZ.md §2):
  * main export   - "Список_всех_пациентов_по_клинике_за_весь_период.xlsx"
  * district files - same format, subsets of the main export split by district
  * conditions    - "ПСОРИАЗ, АТОПИК ДЕРМА, ВИТИЛИГО.xls": diagnosis, first/repeat, phone
  * cold base     - "nomer.xlsx": phones only, several sheets
"""

import re
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date
from functools import cache
from pathlib import Path

from app.core.text import InvalidPhoneError, normalize_uz_phone, search_key
from app.modules.patients.constants import DISTRICT_SPELLINGS
from app.modules.patients.models import Gender

# "Абдуллаев Али Год рождения: 01.02.1990, Пол: Мужской, Домашний адрес: Кува ..."
# (the export sometimes omits the space before "Год рождения")
_FIO = re.compile(
    r"^(?P<name>.*?)\s*Год рождения:\s*(?P<dob>[^,]*?)\s*,\s*Пол:\s*(?P<sex>[^,]*?)\s*,"
    r"\s*Домашний адрес:\s*(?P<addr>.*)$",
    re.S,
)
_DOB = re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})$")
_SEX = {"женский": Gender.FEMALE, "мужской": Gender.MALE}
_FOOTER = {"итого", "регистратор:"}

# file-name fragment -> district, for the per-district files
DISTRICT_FILES: dict[str, str] = {
    "Бешарик": "Beshariq",
    "Богдод": "Bog'dod",
    "Бувайда": "Buvayda",
    "Дангара": "Dang'ara",
    "Узбекистон": "O'zbekiston",
    "Учкупик": "Uchko'prik",
    "Учкуприк": "Uchko'prik",
    "Коканд": "Qo'qon shahri",
}


@dataclass
class ParsedPatient:
    source: str  # "main" | "district"
    row: int
    legacy_no: str
    full_name: str
    birth_date: date | None
    gender: Gender
    address: str | None
    district: str | None
    phone: str | None  # E.164, None if missing/invalid
    raw_phone: str
    diagnosis: str | None
    problems: list[str] = field(default_factory=list)


@dataclass
class ParsedCondition:
    sheet: str
    row: int
    diagnosis: str
    visit_type: str | None  # "first" | "repeat"
    phone: str | None
    raw_phone: str
    problems: list[str] = field(default_factory=list)


@dataclass
class ParsedColdPhone:
    sheet: str
    row: int
    phone: str | None
    raw_phone: str


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_phone(raw: str) -> tuple[str | None, str | None]:
    """Returns (e164, problem)."""
    if not raw:
        return None, "phone_missing"
    try:
        return normalize_uz_phone(raw), None
    except InvalidPhoneError:
        return None, "phone_invalid"


def parse_birth_date(raw: str) -> tuple[date | None, str | None]:
    raw = raw.strip()
    if not raw or raw in {".", ". .", ". .."} or not any(ch.isdigit() for ch in raw):
        return None, None
    m = _DOB.match(raw.replace(" ", ""))
    if not m:
        return None, "birth_date_invalid"
    d, mth, y = (int(x) for x in m.groups())
    try:
        value = date(y, mth, d)
    except ValueError:
        return None, "birth_date_invalid"
    if not (date(1900, 1, 1) <= value <= date.today()):
        return None, "birth_date_invalid"
    return value, None


@cache
def _district_patterns() -> list[tuple[re.Pattern[str], str]]:
    pairs = {
        search_key(spelling): district
        for district, spellings in DISTRICT_SPELLINGS.items()
        for spelling in spellings
    }
    ordered = sorted(pairs.items(), key=lambda kv: -len(kv[0]))
    return [(re.compile(rf"\b{re.escape(key)}\b"), district) for key, district in ordered]


def detect_district(address: str) -> str | None:
    """District of the earliest place name found in a free-text address."""
    # "г.Фергана" must stay two words, so punctuation becomes a space before normalizing
    key = search_key(re.sub(r"[.,;:/()\-]", " ", address))
    best: tuple[int, int, str] | None = None
    for pattern, district in _district_patterns():
        m = pattern.search(key)
        if m and (best is None or (m.start(), -len(m.group())) < best[:2]):
            best = (m.start(), -len(m.group()), district)
    return best[2] if best else None


def parse_fio_cell(cell: str) -> tuple[str, date | None, Gender, str | None, list[str]]:
    m = _FIO.match(cell.strip())
    if not m:
        return re.sub(r"\s+", " ", cell).strip(), None, Gender.UNKNOWN, None, ["fio_unparsed"]
    problems: list[str] = []
    name = re.sub(r"\s+", " ", m["name"]).strip()
    dob, dob_problem = parse_birth_date(m["dob"])
    if dob_problem:
        problems.append(dob_problem)
    gender = _SEX.get(m["sex"].strip().lower(), Gender.UNKNOWN)
    address = re.sub(r"\s+", " ", m["addr"]).strip() or None
    return name, dob, gender, address, problems


def read_patient_list(
    path: Path, source: str, district: str | None = None
) -> Iterator[ParsedPatient]:
    """Rows of the main export or a per-district file (same layout: №, ФИО, Диагноз, Телефон)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        header_seen = False
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = [_cell_text(v) for v in (row + (None,) * 5)[:5]]
            if not header_seen:
                header_seen = cells[1] == "ФИО Пациента"
                continue
            no, fio, diagnosis, raw_phone, _ = cells
            if not fio or fio.lower() in _FOOTER or no.lower() in _FOOTER:
                continue
            name, dob, gender, address, problems = parse_fio_cell(fio)
            phone, phone_problem = parse_phone(raw_phone)
            if phone_problem:
                problems.append(phone_problem)
            if len(name) < 2:
                problems.append("name_missing")
            yield ParsedPatient(
                source=source,
                row=idx,
                legacy_no=no,
                full_name=name,
                birth_date=dob,
                gender=gender,
                address=address,
                district=district or (detect_district(address) if address else None),
                phone=phone,
                raw_phone=raw_phone,
                diagnosis=diagnosis or None,
                problems=problems,
            )
    finally:
        wb.close()


def _sheets(path: Path) -> Iterator[tuple[str, list[list[str]]]]:
    """(sheet name, rows as text) for both legacy .xls and .xlsx workbooks."""
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        for sheet in book.sheets():
            yield (
                sheet.name,
                [
                    [_cell_text(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
                    for r in range(sheet.nrows)
                ],
            )
        return
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            yield ws.title, [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def read_conditions(path: Path) -> Iterator[ParsedCondition]:
    for sheet_name, rows in _sheets(path):
        for r, cells in enumerate(rows[1:], start=2):  # row 1 is the header
            diagnosis, visit, raw_phone = (cells + ["", "", ""])[:3]
            if not diagnosis and not raw_phone:
                continue
            phone, problem = parse_phone(raw_phone)
            visit_type = {"первично": "first", "повторно": "repeat"}.get(visit.lower())
            yield ParsedCondition(
                sheet=sheet_name,
                row=r,
                diagnosis=diagnosis,
                visit_type=visit_type,
                phone=phone,
                raw_phone=raw_phone,
                problems=[problem] if problem else [],
            )


def read_cold_phones(path: Path) -> Iterator[ParsedColdPhone]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                raw = _cell_text(row[0] if row else None)
                if not raw:
                    continue
                phone, _ = parse_phone(raw)
                yield ParsedColdPhone(sheet=ws.title, row=idx, phone=phone, raw_phone=raw)
    finally:
        wb.close()


def district_from_filename(name: str) -> str | None:
    return next((d for fragment, d in DISTRICT_FILES.items() if fragment in name), None)
