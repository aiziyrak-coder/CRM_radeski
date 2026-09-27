"""Legacy Excel import. All data here is synthetic; files mimic the real export layout."""

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select

from app.core.db import SessionLocal
from app.importer.legacy import CHECK_TAG, NO_NAME_TAG, run_import
from app.importer.parsing import detect_district, parse_birth_date, parse_fio_cell
from app.modules.patients.constants import UNKNOWN_NAME
from app.modules.patients.models import Gender, Patient, PatientCondition, PatientKind, PatientPhone


def fio(name: str, dob: str, sex: str, addr: str, gap: str = " ") -> str:
    return f"{name}{gap}Год рождения: {dob}, Пол: {sex}, Домашний адрес: {addr} "


@pytest.mark.parametrize(
    ("cell", "expected"),
    [
        (
            fio("Тестова Мадина", "05.06.1990", "Женский", "Кува ул Намуна 3"),
            ("Тестова Мадина", date(1990, 6, 5), Gender.FEMALE, "Кува ул Намуна 3"),
        ),
        (  # the export sometimes glues the name to "Год рождения"
            fio("Сидоров Иван Петрович", "01.01.1980", "мужской", "Ташкент", gap=""),
            ("Сидоров Иван Петрович", date(1980, 1, 1), Gender.MALE, "Ташкент"),
        ),
        (
            fio("Номаълум Шахс", ". .", "", "г.Фергана"),
            ("Номаълум Шахс", None, Gender.UNKNOWN, "г.Фергана"),
        ),
    ],
)
def test_parse_fio_cell(cell: str, expected: tuple) -> None:
    name, dob, gender, address, problems = parse_fio_cell(cell)
    assert (name, dob, gender, address) == expected
    assert problems == []


@pytest.mark.parametrize("raw", ["01.01.6989", "12.12. 201", "31.02.2000", "1.1.101"])
def test_implausible_birth_dates_are_reported(raw: str) -> None:
    assert parse_birth_date(raw) == (None, "birth_date_invalid")


@pytest.mark.parametrize(
    ("address", "district"),
    [
        ("г.Фергана ул Намуна 1", "Farg'ona shahri"),
        ("Ферганский р-н, Водил", "Farg'ona tumani"),
        ("Кувасой ш", "Quvasoy shahri"),
        ("Кува, ул. Фергана 5", "Quva"),  # earliest place name wins
        ("Ёзёвон тумани", "Yozyovon"),
        ("Қўқон шахар", "Qo'qon shahri"),
        ("Олтарик", "Oltiariq"),
        ("Andijon, Asaka", "Andijon viloyati"),
        ("Ферганская обл.", None),
        ("", None),
    ],
)
def test_detect_district(address: str, district: str | None) -> None:
    assert detect_district(address) == district


# --- end-to-end -------------------------------------------------------------------------------


def _patient_list(path: Path, title: str, rows: list[tuple]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append([None] * 5)
    ws.append([None, title])
    ws.append([None] * 5)
    ws.append(["№", "ФИО Пациента", "Диагноз", "Телефон", "Примечание"])
    for row in rows:
        ws.append(list(row) + [None])
    ws.append([None, None, "ИТОГО", "ИТОГО"])
    ws.append(["Регистратор:", None, "______"])
    wb.save(path)


@pytest.fixture
def data_dir(tmp_path: Path) -> Path:
    _patient_list(
        tmp_path / "Список_всех_пациентов_по_клинике_за_весь_период.xlsx",
        "СПИСОК ПАЦИЕНТОВ",
        [
            (
                1,
                fio("Тестова Мадина", "05.06.1990", "Женский", "Кува ул Намуна 3"),
                "Акне",
                901110001,
            ),
            (2, fio("Синов Бобур", "02.03.1985", "Мужской", "г.Фергана"), "Витилиго", 901110002),
            # same person listed twice with a second phone
            (3, fio("Тестова Мадина", "05.06.1990", "Женский", "Кува"), None, 901110003),
            # broken phone and birth date: imported, but flagged for the operator
            (4, fio("Хатоев Жасур", "01.01.6989", "Мужской", "Бешарик"), "обс", 16685656),
        ],
    )
    _patient_list(
        tmp_path / "Список_всех_пациентов_по_клинике_Бешарик.xlsx",
        "Бешарик",
        [(1, fio("Синов Бобур", "02.03.1985", "Мужской", "г.Фергана"), "Витилиго", 901110002)],
    )

    cond = Workbook()
    ws = cond.active
    ws.title = "витилиго"
    ws.append(["Диагноз:", "Тип приём", "Контактный данные"])
    ws.append(["Витилиго чегараланган", "Повторно", "90-111-00-02"])  # matches Синов Бобур
    ws.append(["Витилиго", "Первично", "90-222-00-01"])  # unknown -> nameless patient
    cond.save(tmp_path / "ПСОРИАЗ, АТОПИК ДЕРМА, ВИТИЛИГО.xlsx")

    cold = Workbook()
    ws = cold.active
    for number in ["90-333-00-01", "90-333-00-01", "90-111-00-01", "12-345", "90-333-00-02"]:
        ws.append([number])
    cold.create_sheet("Лист2").append(["90-333-00-03"])
    cold.save(tmp_path / "nomer.xlsx")
    return tmp_path


async def _count(model) -> int:
    async with SessionLocal() as s:
        return await s.scalar(select(func.count()).select_from(model))


async def test_full_import(data_dir: Path) -> None:
    async with SessionLocal() as session:
        report = await run_import(session, data_dir)

    c = report.counts
    assert c["main:created"] == 3
    assert c["main:same_person_merged"] == 1
    assert c["district:confirmed"] == 1
    assert c["conditions:matched_patient"] == 1
    assert c["conditions:created_without_name"] == 1
    assert c["cold:created"] == 3
    assert c["cold:already_known"] == 2  # duplicate row + a number of a named patient
    assert c["cold:invalid_phone"] == 1
    assert {p["problem"] for p in report.problems} == {"phone_invalid", "birth_date_invalid"}

    async with SessionLocal() as s:
        madina = await s.scalar(select(Patient).where(Patient.full_name == "Тестова Мадина"))
        assert madina.kind == PatientKind.LEGACY
        assert madina.district == "Quva"
        assert madina.gender == Gender.FEMALE
        assert sorted(p.number for p in madina.phones) == ["+998901110001", "+998901110003"]

        bobur = await s.scalar(select(Patient).where(Patient.full_name == "Синов Бобур"))
        assert bobur.district == "Beshariq"  # district file overrides the address guess
        conditions = await s.scalars(
            select(PatientCondition.raw_text).where(PatientCondition.patient_id == bobur.id)
        )
        assert sorted(conditions) == ["Витилиго", "Витилиго чегараланган"]

        broken = await s.scalar(select(Patient).where(Patient.full_name == "Хатоев Жасур"))
        assert broken.phones == [] and broken.birth_date is None
        assert CHECK_TAG in broken.tags and "16685656" in broken.notes

        nameless = await s.scalar(select(Patient).where(Patient.tags.contains([NO_NAME_TAG])))
        assert nameless.full_name == UNKNOWN_NAME

        cold = await s.scalar(
            select(func.count()).select_from(Patient).where(Patient.kind == PatientKind.COLD)
        )
        assert cold == 3


async def test_import_is_idempotent(data_dir: Path) -> None:
    async with SessionLocal() as session:
        await run_import(session, data_dir)
    before = (await _count(Patient), await _count(PatientPhone), await _count(PatientCondition))

    async with SessionLocal() as session:
        report = await run_import(session, data_dir)

    after = (await _count(Patient), await _count(PatientPhone), await _count(PatientCondition))
    assert after == before
    assert report.counts["main:already_imported"] == 3
    assert report.counts["main:same_person_merged"] == 1  # the repeated row, matched again
    assert report.counts["cold:created"] == 0


async def test_dry_run_writes_nothing(data_dir: Path) -> None:
    async with SessionLocal() as session:
        report = await run_import(session, data_dir, dry_run=True)

    assert report.counts["main:created"] == 3
    assert await _count(Patient) == 0


async def test_reimport_after_a_merge_brings_no_duplicate(data_dir: Path) -> None:
    from app.modules.patients import service as patients

    async with SessionLocal() as session:
        await run_import(session, data_dir)
    async with SessionLocal() as s:
        # an operator folds the imported card into one made in the CRM
        imported = await s.scalar(select(Patient).where(Patient.full_name == "Синов Бобур"))
        crm_card = Patient(
            full_name="Бобур Синов", search_key="bobur sinov", kind=PatientKind.ACTIVE
        )
        s.add(crm_card)
        await s.commit()
        target, source = await patients.lock_pair(s, crm_card.id, imported.id)
        await patients.merge(s, target, source)
        await s.commit()
    before = await _count(Patient)

    async with SessionLocal() as session:
        report = await run_import(session, data_dir)
    assert await _count(Patient) == before
    assert report.counts["district:created"] == 0 and report.counts["main:created"] == 0
